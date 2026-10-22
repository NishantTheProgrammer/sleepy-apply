from fastapi import APIRouter, HTTPException, Request, UploadFile, File
from bson import ObjectId
from app.db.client import get_db
from app.security.jwt import get_current_user_id
from app.models.profile import Profile, ProfileQuestion
from datetime import datetime, timezone
from pydantic import BaseModel
import structlog

router = APIRouter(prefix="/profile", tags=["profile"])
logger = structlog.get_logger()


@router.get("")
async def get_profile(request: Request):
    user_id = get_current_user_id(request)
    db = get_db()
    doc = await db.profiles.find_one({"user_id": ObjectId(user_id)})
    if not doc:
        return {"profile": None}
    doc["_id"] = str(doc["_id"])
    doc["user_id"] = str(doc["user_id"])
    return {"profile": doc}


@router.patch("")
async def update_profile(request: Request):
    """Partial update of the user's profile."""
    user_id = get_current_user_id(request)
    db = get_db()
    body = await request.json()
    now = datetime.now(timezone.utc)

    # Whitelist allowed top-level keys
    allowed = {
        "personal", "summary", "skills", "experience", "education",
        "projects", "certifications", "job_preferences", "jobPreferences",
        "achievements", "additionalInformation", "additional_information",
        "selected_template",
    }
    update = {k: v for k, v in body.items() if k in allowed}
    if not update:
        raise HTTPException(status_code=400, detail="No valid fields to update")

    # Recalculate completion score
    from app.agents.profile_creator import ProfileCreatorAgent
    agent = ProfileCreatorAgent()
    current_doc = await db.profiles.find_one({"user_id": ObjectId(user_id)}) or {}
    merged = {**current_doc, **update}
    update["completion_score"] = agent._calculate_completion_score(merged)
    update["updated_at"] = now

    await db.profiles.update_one(
        {"user_id": ObjectId(user_id)},
        {"$set": update, "$setOnInsert": {"user_id": ObjectId(user_id), "created_at": now}},
        upsert=True,
    )
    doc = await db.profiles.find_one({"user_id": ObjectId(user_id)})
    doc["_id"] = str(doc["_id"])
    doc["user_id"] = str(doc["user_id"])
    return {"profile": doc}


@router.post("/resume")
async def upload_resume(request: Request, file: UploadFile = File(...)):
    """Upload a resume (PDF or DOCX) — triggers the Profile Creator Agent."""
    user_id = get_current_user_id(request)

    filename_lower = (file.filename or "").lower()
    if not (filename_lower.endswith(".pdf") or filename_lower.endswith(".docx")):
        raise HTTPException(status_code=400, detail="Only PDF and DOCX files are accepted")

    content = await file.read()
    if len(content) > 10 * 1024 * 1024:  # 10MB limit
        raise HTTPException(status_code=400, detail="File too large (max 10MB)")

    text = ""
    import io
    if filename_lower.endswith(".pdf"):
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(content))
            # Extract text preserving visual column layout
            pages_text = []
            for page in reader.pages:
                try:
                    p_txt = page.extract_text(extraction_mode="layout")
                except Exception:
                    p_txt = page.extract_text()
                if p_txt:
                    pages_text.append(p_txt)
            text = "\n".join(pages_text)
            if not text.strip():
                # Fallback to plain extraction
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as e:
            logger.error("Failed to extract PDF text", error=str(e), filename=file.filename)
            raise HTTPException(status_code=400, detail=f"Failed to read PDF: {str(e)}")
    elif filename_lower.endswith(".docx"):
        import docx
        try:
            doc = docx.Document(io.BytesIO(content))
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            table_cells = [cell.text for t in doc.tables for row in t.rows for cell in row.cells if cell.text.strip()]
            text = "\n".join(paras + table_cells)
        except Exception as e:
            logger.error("Failed to extract DOCX text", error=str(e), filename=file.filename)
            raise HTTPException(status_code=400, detail=f"Failed to read DOCX: {str(e)}")

    if not text.strip():
        raise HTTPException(status_code=400, detail="Could not extract any readable text from uploaded file")

    # Clean excessive whitespace while preserving section layout
    import re
    cleaned_text = re.sub(r"\n{3,}", "\n\n", text).strip()

    # Store raw text and set is_analyzing flag
    db = get_db()
    now = datetime.now(timezone.utc)
    await db.profiles.update_one(
        {"user_id": ObjectId(user_id)},
        {
            "$set": {
                "resume_filename": file.filename,
                "resume_raw_text": cleaned_text,
                "is_analyzing": True,
                "analysis_started_at": now,
                "analysis_error": None,
                "updated_at": now,
            },
            "$setOnInsert": {"user_id": ObjectId(user_id), "created_at": now},
        },
        upsert=True,
    )

    import asyncio
    from app.graphs.profile_graph import profile_pipeline
    asyncio.create_task(profile_pipeline.ainvoke({"user_id": user_id, "resume_text": cleaned_text}))
    logger.info("Resume uploaded and profile pipeline triggered", user_id=user_id, filename=file.filename, chars=len(cleaned_text))

    return {
        "ok": True,
        "is_analyzing": True,
        "message": "Resume uploaded. AI is analyzing your profile...",
        "filename": file.filename,
    }


@router.get("/questions")
async def get_pending_questions(request: Request):
    """Return AI-generated questions awaiting user answers."""
    user_id = get_current_user_id(request)
    db = get_db()
    doc = await db.profiles.find_one(
        {"user_id": ObjectId(user_id)},
        {"pending_questions": 1},
    )
    if not doc:
        return {"questions": []}
    questions = [q for q in doc.get("pending_questions", []) if not q.get("answered_at")]
    return {"questions": questions}


class AnswerRequest(BaseModel):
    question_id: str
    answer: str


@router.post("/answers")
async def submit_answers(request: Request, body: list[AnswerRequest]):
    """Submit user answers to AI profile questions."""
    user_id = get_current_user_id(request)
    db = get_db()
    now = datetime.now(timezone.utc)

    for item in body:
        await db.profiles.update_one(
            {
                "user_id": ObjectId(user_id),
                "pending_questions.id": item.question_id,
            },
            {
                "$set": {
                    "pending_questions.$.answer": item.answer,
                    "pending_questions.$.answered_at": now,
                    "updated_at": now,
                }
            },
        )

    # TODO: trigger profile_graph to incorporate answers
    return {"ok": True, "answered": len(body)}
