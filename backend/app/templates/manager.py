import os
import io
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docxtpl import DocxTemplate
import structlog

logger = structlog.get_logger()

TEMPLATES_DIR = Path(__file__).parent / "files"
TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)

TEMPLATE_CATALOG = [
    {
        "id": "default",
        "name": "Default Clean ATS",
        "tagline": "Custom layout matching your Google Doc with dual-side header & section dividers",
        "badge": "Default",
        "badge_color": "emerald",
        "ats_score": 99,
        "style": "Clean Dual-Header",
        "primary_color": "#1e293b",
        "secondary_color": "#0f172a",
        "font_family": "Calibri / Arial",
        "layout": "Dual-Side Header + Dividers",
        "ideal_for": "Full Stack & Software Engineering Roles",
        "features": [
            "Google Doc Exact Match",
            "Inline Contact Bar & Links",
            "Compact Skills Grid",
            "High ATS Readability (99%)"
        ],
    },
    {
        "id": "modern_tech",
        "name": "Modern Tech ATS",
        "tagline": "Clean modern layout optimized for Software Engineers & Tech Leads",
        "badge": "Popular",
        "badge_color": "indigo",
        "ats_score": 98,
        "style": "Modern Minimalist",
        "primary_color": "#4f46e5",
        "secondary_color": "#0f172a",
        "font_family": "Calibri / Inter",
        "layout": "Clean Single Column",
        "ideal_for": "Full Stack, Backend, Frontend & DevOps Engineers",
        "features": [
            "Category-based Skills Matrix",
            "Project & Technology Tagging",
            "High ATS Readability (98%)",
            "Metrics & Impact Focus"
        ],
    },
    {
        "id": "minimal_executive",
        "name": "Minimal Executive",
        "tagline": "Elegant, serif-accented design for Senior Engineers & Architects",
        "badge": "Executive",
        "badge_color": "emerald",
        "ats_score": 96,
        "style": "Executive Classic",
        "primary_color": "#0f766e",
        "secondary_color": "#1e293b",
        "font_family": "Georgia / Garamond",
        "layout": "Sophisticated Flow",
        "ideal_for": "Engineering Managers, Architects & Senior Tech Roles",
        "features": [
            "Leadership & Scope Emphasis",
            "Refined Typography",
            "Impact Metrics Highlighting",
            "Clean Section Dividers"
        ],
    },
    {
        "id": "compact_developer",
        "name": "Compact Developer",
        "tagline": "Dense, tech-forward format emphasizing hands-on coding & projects",
        "badge": "Compact",
        "badge_color": "purple",
        "ats_score": 97,
        "style": "Developer Dense",
        "primary_color": "#7c3aed",
        "secondary_color": "#18181b",
        "font_family": "Arial / Consolas",
        "layout": "Compact Single Page",
        "ideal_for": "Hands-on Developers, Open Source Contributors & Hackers",
        "features": [
            "Maximized Page Space",
            "Inline Tech Stack Badges",
            "GitHub & Project Centric",
            "Bullet Point Density"
        ],
    },
    {
        "id": "classic_ats",
        "name": "Classic Ivy ATS",
        "tagline": "Harvard-standard monochrome layout with 100% ATS guarantee",
        "badge": "100% ATS",
        "badge_color": "slate",
        "ats_score": 100,
        "style": "Traditional Monochrome",
        "primary_color": "#000000",
        "secondary_color": "#334155",
        "font_family": "Times New Roman",
        "layout": "Universal Standard",
        "ideal_for": "Enterprise, Corporate & Fortune 500 Applications",
        "features": [
            "100% ATS Parse Rate",
            "Zero Styling Clutter",
            "Strict Chronological Order",
            "Accepted Worldwide"
        ],
    },
]


def _add_bottom_border(paragraph, color_hex="CCCCCC", sz="4"):
    """Add a subtle bottom border underline to a section heading paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), sz)
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color_hex.replace("#", ""))
    pBdr.append(bottom)
    pPr.append(pBdr)


def create_modern_tech_template(output_path: Path):
    """Creates Modern Tech ATS .docx template with Jinja2 placeholders."""
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    # 1. Header: Name
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    r_name = p_name.add_run("{{ personal.fullName }}")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(20)
    r_name.font.bold = True
    r_name.font.color.rgb = RGBColor(15, 23, 42)  # #0f172a

    # Contact line
    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_after = Pt(8)
    r_contact = p_contact.add_run("{{ personal.contact_line }}")
    r_contact.font.name = "Calibri"
    r_contact.font.size = Pt(9.5)
    r_contact.font.color.rgb = RGBColor(71, 85, 105)

    # 2. Professional Summary
    p_sum_head = doc.add_paragraph()
    p_sum_head.paragraph_format.space_before = Pt(6)
    p_sum_head.paragraph_format.space_after = Pt(3)
    r_sum_head = p_sum_head.add_run("PROFESSIONAL SUMMARY")
    r_sum_head.font.name = "Calibri"
    r_sum_head.font.size = Pt(11)
    r_sum_head.font.bold = True
    r_sum_head.font.color.rgb = RGBColor(79, 70, 229)  # Indigo
    _add_bottom_border(p_sum_head, "4F46E5", "8")

    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_after = Pt(6)
    r_sum = p_sum.add_run("{{ summary }}")
    r_sum.font.name = "Calibri"
    r_sum.font.size = Pt(9.5)
    r_sum.font.color.rgb = RGBColor(51, 65, 85)

    # 3. Technical Skills
    p_sk_head = doc.add_paragraph()
    p_sk_head.paragraph_format.space_before = Pt(6)
    p_sk_head.paragraph_format.space_after = Pt(3)
    r_sk_head = p_sk_head.add_run("TECHNICAL SKILLS")
    r_sk_head.font.name = "Calibri"
    r_sk_head.font.size = Pt(11)
    r_sk_head.font.bold = True
    r_sk_head.font.color.rgb = RGBColor(79, 70, 229)
    _add_bottom_border(p_sk_head, "4F46E5", "8")

    p_sk_loop = doc.add_paragraph()
    p_sk_loop.paragraph_format.space_after = Pt(2)
    r_sk_loop = p_sk_loop.add_run("{% for cat in skills_categories %}• {{ cat.name }}: {{ cat.items_str }}\n{% endfor %}")
    r_sk_loop.font.name = "Calibri"
    r_sk_loop.font.size = Pt(9.5)
    r_sk_loop.font.color.rgb = RGBColor(51, 65, 85)

    # 4. Work Experience
    p_exp_head = doc.add_paragraph()
    p_exp_head.paragraph_format.space_before = Pt(8)
    p_exp_head.paragraph_format.space_after = Pt(3)
    r_exp_head = p_exp_head.add_run("WORK EXPERIENCE")
    r_exp_head.font.name = "Calibri"
    r_exp_head.font.size = Pt(11)
    r_exp_head.font.bold = True
    r_exp_head.font.color.rgb = RGBColor(79, 70, 229)
    _add_bottom_border(p_exp_head, "4F46E5", "8")

    p_exp_loop = doc.add_paragraph()
    p_exp_loop.paragraph_format.space_after = Pt(4)
    r_exp_loop = p_exp_loop.add_run(
        "{% for exp in experience %}"
        "{{ exp.title }} — {{ exp.company }} ({{ exp.dates }})\n"
        "{% if exp.description %}{{ exp.description }}\n{% endif %}"
        "{% for resp in exp.responsibilities %}  • {{ resp }}\n{% endfor %}"
        "{% for ach in exp.achievements %}  • {{ ach }}\n{% endfor %}\n"
        "{% endfor %}"
    )
    r_exp_loop.font.name = "Calibri"
    r_exp_loop.font.size = Pt(9.5)
    r_exp_loop.font.color.rgb = RGBColor(51, 65, 85)

    # 5. Projects
    p_proj_head = doc.add_paragraph()
    p_proj_head.paragraph_format.space_before = Pt(6)
    p_proj_head.paragraph_format.space_after = Pt(3)
    r_proj_head = p_proj_head.add_run("KEY PROJECTS")
    r_proj_head.font.name = "Calibri"
    r_proj_head.font.size = Pt(11)
    r_proj_head.font.bold = True
    r_proj_head.font.color.rgb = RGBColor(79, 70, 229)
    _add_bottom_border(p_proj_head, "4F46E5", "8")

    p_proj_loop = doc.add_paragraph()
    p_proj_loop.paragraph_format.space_after = Pt(4)
    r_proj_loop = p_proj_loop.add_run(
        "{% for proj in projects %}"
        "{{ proj.name }} | {{ proj.tech_str }}\n"
        "{% if proj.description %}  • {{ proj.description }}\n{% endif %}\n"
        "{% endfor %}"
    )
    r_proj_loop.font.name = "Calibri"
    r_proj_loop.font.size = Pt(9.5)
    r_proj_loop.font.color.rgb = RGBColor(51, 65, 85)

    # 6. Education
    p_edu_head = doc.add_paragraph()
    p_edu_head.paragraph_format.space_before = Pt(6)
    p_edu_head.paragraph_format.space_after = Pt(3)
    r_edu_head = p_edu_head.add_run("EDUCATION")
    r_edu_head.font.name = "Calibri"
    r_edu_head.font.size = Pt(11)
    r_edu_head.font.bold = True
    r_edu_head.font.color.rgb = RGBColor(79, 70, 229)
    _add_bottom_border(p_edu_head, "4F46E5", "8")

    p_edu_loop = doc.add_paragraph()
    p_edu_loop.paragraph_format.space_after = Pt(2)
    r_edu_loop = p_edu_loop.add_run(
        "{% for edu in education %}"
        "{{ edu.degree }} — {{ edu.institution }} ({{ edu.year }})\n"
        "{% endfor %}"
    )
    r_edu_loop.font.name = "Calibri"
    r_edu_loop.font.size = Pt(9.5)
    r_edu_loop.font.color.rgb = RGBColor(51, 65, 85)

    doc.save(str(output_path))
    logger.info("Created template file", path=str(output_path))


def create_minimal_executive_template(output_path: Path):
    """Creates Minimal Executive .docx template with classic styling."""
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)

    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_name = p_name.add_run("{{ personal.fullName }}")
    r_name.font.name = "Georgia"
    r_name.font.size = Pt(22)
    r_name.font.bold = True
    r_name.font.color.rgb = RGBColor(15, 118, 110)  # Teal #0f766e

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_after = Pt(10)
    r_contact = p_contact.add_run("{{ personal.contact_line }}")
    r_contact.font.name = "Georgia"
    r_contact.font.size = Pt(9)
    r_contact.font.color.rgb = RGBColor(100, 116, 139)

    for section_title, content_code in [
        ("EXECUTIVE SUMMARY", "{{ summary }}"),
        ("CORE COMPETENCIES", "{% for cat in skills_categories %}• {{ cat.name }}: {{ cat.items_str }}\n{% endfor %}"),
        ("PROFESSIONAL EXPERIENCE", "{% for exp in experience %}{{ exp.title }} | {{ exp.company }} | {{ exp.dates }}\n{% for resp in exp.responsibilities %}  • {{ resp }}\n{% endfor %}\n{% endfor %}"),
        ("NOTABLE PROJECTS", "{% for proj in projects %}{{ proj.name }} ({{ proj.tech_str }}): {{ proj.description }}\n{% endfor %}"),
        ("ACADEMIC CREDENTIALS", "{% for edu in education %}{{ edu.degree }} | {{ edu.institution }} | {{ edu.year }}\n{% endfor %}"),
    ]:
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(8)
        p_head.paragraph_format.space_after = Pt(2)
        r_head = p_head.add_run(section_title)
        r_head.font.name = "Georgia"
        r_head.font.size = Pt(10.5)
        r_head.font.bold = True
        r_head.font.color.rgb = RGBColor(15, 118, 110)
        _add_bottom_border(p_head, "0F766E", "6")

        p_body = doc.add_paragraph()
        p_body.paragraph_format.space_after = Pt(4)
        r_body = p_body.add_run(content_code)
        r_body.font.name = "Georgia"
        r_body.font.size = Pt(9.5)
        r_body.font.color.rgb = RGBColor(30, 41, 59)

    doc.save(str(output_path))


def create_compact_developer_template(output_path: Path):
    """Creates Compact Developer .docx template with high density."""
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.4)
        s.bottom_margin = Inches(0.4)
        s.left_margin = Inches(0.5)
        s.right_margin = Inches(0.5)

    p_name = doc.add_paragraph()
    r_name = p_name.add_run("{{ personal.fullName }}")
    r_name.font.name = "Arial"
    r_name.font.size = Pt(18)
    r_name.font.bold = True
    r_name.font.color.rgb = RGBColor(124, 58, 237)  # Purple

    p_contact = doc.add_paragraph()
    p_contact.paragraph_format.space_after = Pt(6)
    r_contact = p_contact.add_run("{{ personal.contact_line }}")
    r_contact.font.name = "Arial"
    r_contact.font.size = Pt(8.5)
    r_contact.font.color.rgb = RGBColor(71, 85, 105)

    for section_title, content_code in [
        ("SUMMARY", "{{ summary }}"),
        ("SKILLS", "{% for cat in skills_categories %}[{{ cat.name }}] {{ cat.items_str }}\n{% endfor %}"),
        ("EXPERIENCE", "{% for exp in experience %}> {{ exp.title }} @ {{ exp.company }} ({{ exp.dates }})\n{% for resp in exp.responsibilities %}  * {{ resp }}\n{% endfor %}\n{% endfor %}"),
        ("PROJECTS", "{% for proj in projects %}> {{ proj.name }} [{{ proj.tech_str }}] - {{ proj.description }}\n{% endfor %}"),
        ("EDUCATION", "{% for edu in education %}> {{ edu.degree }} - {{ edu.institution }} ({{ edu.year }})\n{% endfor %}"),
    ]:
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(5)
        p_head.paragraph_format.space_after = Pt(2)
        r_head = p_head.add_run(section_title)
        r_head.font.name = "Arial"
        r_head.font.size = Pt(10)
        r_head.font.bold = True
        r_head.font.color.rgb = RGBColor(124, 58, 237)
        _add_bottom_border(p_head, "7C3AED", "6")

        p_body = doc.add_paragraph()
        p_body.paragraph_format.space_after = Pt(3)
        r_body = p_body.add_run(content_code)
        r_body.font.name = "Arial"
        r_body.font.size = Pt(8.5)
        r_body.font.color.rgb = RGBColor(24, 24, 27)

    doc.save(str(output_path))


def create_classic_ats_template(output_path: Path):
    """Creates Classic Ivy ATS .docx template in 100% black-and-white standard."""
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.5)
        s.right_margin = Inches(0.5)

    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_name = p_name.add_run("{{ personal.fullName }}")
    r_name.font.name = "Times New Roman"
    r_name.font.size = Pt(18)
    r_name.font.bold = True
    r_name.font.color.rgb = RGBColor(0, 0, 0)

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_after = Pt(8)
    r_contact = p_contact.add_run("{{ personal.contact_line }}")
    r_contact.font.name = "Times New Roman"
    r_contact.font.size = Pt(9.5)
    r_contact.font.color.rgb = RGBColor(0, 0, 0)

    for section_title, content_code in [
        ("SUMMARY", "{{ summary }}"),
        ("TECHNICAL SKILLS", "{% for cat in skills_categories %}{{ cat.name }}: {{ cat.items_str }}\n{% endfor %}"),
        ("PROFESSIONAL EXPERIENCE", "{% for exp in experience %}{{ exp.company }} — {{ exp.title }} ({{ exp.dates }})\n{% for resp in exp.responsibilities %}• {{ resp }}\n{% endfor %}\n{% endfor %}"),
        ("PROJECTS", "{% for proj in projects %}{{ proj.name }} | {{ proj.tech_str }}: {{ proj.description }}\n{% endfor %}"),
        ("EDUCATION", "{% for edu in education %}{{ edu.institution }} — {{ edu.degree }} ({{ edu.year }})\n{% endfor %}"),
    ]:
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(6)
        p_head.paragraph_format.space_after = Pt(2)
        r_head = p_head.add_run(section_title)
        r_head.font.name = "Times New Roman"
        r_head.font.size = Pt(10)
        r_head.font.bold = True
        r_head.font.color.rgb = RGBColor(0, 0, 0)
        _add_bottom_border(p_head, "000000", "4")

        p_body = doc.add_paragraph()
        p_body.paragraph_format.space_after = Pt(4)
        r_body = p_body.add_run(content_code)
        r_body.font.name = "Times New Roman"
        r_body.font.size = Pt(9.5)
        r_body.font.color.rgb = RGBColor(0, 0, 0)

    doc.save(str(output_path))


def ensure_templates_exist():
    """Ensure all template .docx files exist on disk."""
    creators = {
        "modern_tech": create_modern_tech_template,
        "minimal_executive": create_minimal_executive_template,
        "compact_developer": create_compact_developer_template,
        "classic_ats": create_classic_ats_template,
    }
    for t_id, creator in creators.items():
        path = TEMPLATES_DIR / f"{t_id}.docx"
        if not path.exists():
            creator(path)

    path_default = TEMPLATES_DIR / "default.docx"
    if not path_default.exists() and (TEMPLATES_DIR / "default_backup.docx").exists():
        import shutil
        shutil.copyfile(TEMPLATES_DIR / "default_backup.docx", path_default)


def format_profile_for_rendering(profile: dict) -> dict:
    """Transform candidate profile dictionary into standard template rendering context."""
    personal = profile.get("personal", {})
    name = personal.get("fullName") or personal.get("full_name") or "Candidate Name"
    email = personal.get("email", "")
    phone = personal.get("phone", "")

    # Location string
    loc = personal.get("location")
    loc_str = ""
    if isinstance(loc, dict):
        loc_parts = [loc.get(k) for k in ["city", "country"] if loc.get(k)]
        loc_str = ", ".join(loc_parts)
    elif isinstance(loc, str):
        loc_str = loc

    # Contact line parts
    contact_parts = []
    if email: contact_parts.append(email)
    if phone: contact_parts.append(phone)
    if loc_str: contact_parts.append(loc_str)
    for k in ["linkedin", "github", "portfolio"]:
        val = personal.get(k)
        if val:
            clean = str(val).replace("https://", "").replace("http://", "")
            contact_parts.append(clean)

    contact_line = " • ".join(contact_parts)

    # Skills categories and lines
    skills_raw = profile.get("skills", {})
    skills_categories = []
    skills_lines = []

    if isinstance(skills_raw, dict):
        for cat, items in skills_raw.items():
            if items:
                skills_categories.append({
                    "name": cat.capitalize(),
                    "items": items,
                    "items_str": ", ".join(items)
                })
                for i in range(0, len(items), 3):
                    skills_lines.append(", ".join(items[i:i + 3]))
    elif isinstance(skills_raw, list):
        skills_categories.append({
            "name": "Technical Skills",
            "items": skills_raw,
            "items_str": ", ".join(skills_raw)
        })
        for i in range(0, len(skills_raw), 3):
            skills_lines.append(", ".join(skills_raw[i:i + 3]))

    # Experience
    raw_exp = profile.get("experience", [])
    experience = []
    for exp in raw_exp:
        title = exp.get("jobTitle") or exp.get("title", "")
        company = exp.get("company", "")
        start = exp.get("startDate") or exp.get("start_date", "")
        is_curr = exp.get("current") or exp.get("is_current", False)
        end = "Present" if is_curr else (exp.get("endDate") or exp.get("end_date", ""))
        dates = f"{start} – {end}".strip(" –")

        resps = exp.get("responsibilities", [])
        achs = exp.get("achievements", [])

        # Sub-projects achievements and formatted project items
        projects_formatted = []
        for sub in exp.get("projects", []):
            for ach in sub.get("achievements", []):
                achs.append(f"[{sub.get('name', 'Project')}] {ach}")
            p_desc = sub.get("description", "")
            if not p_desc and sub.get("achievements"):
                p_desc = " ".join(sub.get("achievements"))
            projects_formatted.append({
                "title": sub.get("name", "Project"),
                "description": p_desc
            })

        experience.append({
            "title": title,
            "company": company,
            "dates": dates,
            "description": exp.get("description", ""),
            "responsibilities": resps,
            "achievements": achs,
            "projects_formatted": projects_formatted,
        })

    # Projects
    raw_proj = profile.get("projects", [])
    projects = []
    for p in raw_proj:
        tech = p.get("technologies", [])
        projects.append({
            "name": p.get("name", ""),
            "description": p.get("description", ""),
            "tech": tech,
            "tech_str": ", ".join(tech),
        })

    # Education
    raw_edu = profile.get("education", [])
    education = []
    for ed in raw_edu:
        education.append({
            "degree": ed.get("degree", ""),
            "institution": ed.get("institution", ""),
            "year": ed.get("year") or ed.get("graduation_year", ""),
        })

    has_linkedin = bool(personal.get("linkedin"))
    has_github = bool(personal.get("github"))
    has_portfolio = bool(personal.get("portfolio"))

    return {
        "personal": {
            "fullName": name,
            "email": email,
            "phone": phone,
            "location_str": loc_str,
            "contact_line": contact_line,
            "linkedin": personal.get("linkedin", ""),
            "linkedin_label": "LinkedIn®" if has_linkedin else "",
            "github": personal.get("github", ""),
            "github_label": "GitHub" if has_github else "",
            "portfolio": personal.get("portfolio", ""),
            "portfolio_label": personal.get("portfolio", "") or "portfolio",
        },
        "summary": profile.get("summary", ""),
        "skills_categories": skills_categories,
        "skills_lines": skills_lines,
        "experience": experience,
        "projects": projects,
        "education": education,
        "achievements": profile.get("achievements", []),
    }


def render_docx_for_profile(template_id: str, profile: dict) -> io.BytesIO:
    """Render a .docx resume using docxtpl and the chosen template."""
    ensure_templates_exist()
    template_path = TEMPLATES_DIR / f"{template_id}.docx"
    if not template_path.exists():
        template_path = TEMPLATES_DIR / "default.docx"
    if not template_path.exists():
        template_path = TEMPLATES_DIR / "modern_tech.docx"

    context = format_profile_for_rendering(profile)
    doc = DocxTemplate(template_path)
    doc.render(context)

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream
