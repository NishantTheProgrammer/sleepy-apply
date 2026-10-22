import docx
from pathlib import Path
import docxtpl
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC = Path("app/templates/files/default_backup.docx")
OUT = Path("app/templates/files/default.docx")

doc = docx.Document(SRC)

# Helper functions
def delete_paragraph(p):
    elem = p._p
    elem.getparent().remove(elem)

def insert_paragraph_after(paragraph, text="", style="Normal"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p_obj = docx.text.paragraph.Paragraph(new_p, paragraph._parent)
    try:
        p_obj.style = style
    except Exception:
        pass
    if text:
        p_obj.text = text
    return p_obj

# 1. Update Header P[0]
p0 = doc.paragraphs[0]
for r in p0.runs:
    if "Nishant" in r.text:
        r.text = r.text.replace("Nishant", "{{ personal.fullName }}")
    elif "knishant113@gmail.com" in r.text:
        r.text = r.text.replace("knishant113@gmail.com", "{{ personal.email }}")
    elif "+91 7838716089" in r.text:
        r.text = r.text.replace("+91 7838716089", "{{ personal.phone }}")

for child in p0._p:
    if child.tag.endswith("hyperlink"):
        for r in child.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
            if "LinkedIn" in r.text:
                r.text = "{{ personal.linkedin_label }}"

# 2. Update Header P[1]
p1 = doc.paragraphs[1]
for r in p1.runs:
    if "Ashok Nagar" in r.text:
        prefix, sep, suffix = r.text.partition("\t")
        r.text = "{{ personal.location_str }}" + sep + suffix

for child in p1._p:
    if child.tag.endswith("hyperlink"):
        for r in child.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
            if "GitHub" in r.text:
                r.text = "{{ personal.github_label }}"
            elif "nishanttheprogrammer.com" in r.text:
                r.text = "{{ personal.portfolio_label }}"

# 3. Update Summary P[5]
p5 = doc.paragraphs[5]
p5.text = "{{ summary }}"

# Find section heading indices
def get_p_indices():
    skills_i = exp_i = proj_i = edu_i = -1
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt == "SKILLS":
            skills_i = i
        elif txt == "EXPERIENCE":
            exp_i = i
        elif txt == "RELEVANT PROJECTS":
            proj_i = i
        elif txt == "EDUCATION":
            edu_i = i
    return skills_i, exp_i, proj_i, edu_i

# 4. Update SKILLS: delete static paragraphs between SKILLS and EXPERIENCE
skills_i, exp_i, proj_i, edu_i = get_p_indices()
p_skills = doc.paragraphs[skills_i]

# Delete paragraphs from skills_i + 1 up to exp_i - 1
for _ in range(exp_i - skills_i - 1):
    delete_paragraph(doc.paragraphs[skills_i + 1])

# Insert dynamic skills loop after p_skills
cur = insert_paragraph_after(p_skills, "{%p for sk in skills_lines %}")
p_sk_item = insert_paragraph_after(cur, "{{ sk }}", style="normal")
if p_sk_item.runs:
    p_sk_item.runs[0].font.bold = True
cur = insert_paragraph_after(p_sk_item, "{%p endfor %}")
cur = insert_paragraph_after(cur, "")  # empty spacer before EXPERIENCE

# 5. Update EXPERIENCE
skills_i, exp_i, proj_i, edu_i = get_p_indices()
p_exp = doc.paragraphs[exp_i]

for _ in range(proj_i - exp_i - 1):
    delete_paragraph(doc.paragraphs[exp_i + 1])

cur = insert_paragraph_after(p_exp, "{%p for exp in experience %}")

p_comp = insert_paragraph_after(cur, "{{ exp.company }}", style="Heading 3")
p_dates = insert_paragraph_after(p_comp, "{{ exp.dates }}", style="Heading 3")
p_desc = insert_paragraph_after(p_dates, "{% if exp.description %}{{ exp.description }}{% endif %}", style="normal")
p_sp1 = insert_paragraph_after(p_desc, "", style="normal")

p_title = insert_paragraph_after(p_sp1, "{{ exp.title }}", style="normal")
if p_title.runs:
    p_title.runs[0].font.bold = True

p_resp = insert_paragraph_after(p_title, "{% for resp in exp.responsibilities %}{{ resp }}\n{% endfor %}", style="normal")

p_proj_loop = insert_paragraph_after(p_resp, "{% for p in exp.projects_formatted %}{{ p.title }}: {{ p.description }}\n{% endfor %}", style="normal")

p_sp2 = insert_paragraph_after(p_proj_loop, "", style="normal")
cur = insert_paragraph_after(p_sp2, "{%p endfor %}")
cur = insert_paragraph_after(cur, "")  # spacer

# 6. Update RELEVANT PROJECTS
skills_i, exp_i, proj_i, edu_i = get_p_indices()
p_proj = doc.paragraphs[proj_i]

for _ in range(edu_i - proj_i - 1):
    delete_paragraph(doc.paragraphs[proj_i + 1])

cur = insert_paragraph_after(p_proj, "{%p for proj in projects %}")
p_proj_item = insert_paragraph_after(cur, "{{ proj.name }}: {{ proj.description }}", style="normal")
p_proj_sp = insert_paragraph_after(p_proj_item, "", style="normal")
cur = insert_paragraph_after(p_proj_sp, "{%p endfor %}")
cur = insert_paragraph_after(cur, "")

# 7. Update EDUCATION
skills_i, exp_i, proj_i, edu_i = get_p_indices()
p_edu = doc.paragraphs[edu_i]

# Delete remaining after p_edu
total = len(doc.paragraphs)
for _ in range(total - edu_i - 1):
    delete_paragraph(doc.paragraphs[edu_i + 1])

cur = insert_paragraph_after(p_edu, "{%p for edu in education %}")
p_edu_item = insert_paragraph_after(cur, "{{ edu.degree }}. {{ edu.institution }} ({{ edu.year }})", style="normal")
p_edu_sp = insert_paragraph_after(p_edu_item, "", style="normal")
cur = insert_paragraph_after(p_edu_sp, "{%p endfor %}")

doc.save(OUT)
print("Updated default.docx with dynamic variables successfully!")
